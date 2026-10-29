"""Builds Half_Marathon_Plan_2027.xlsx from plan_data.py and meals.py (all totals as live formulas)."""
import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference

import meals as M
import plan_data as P

OUT = sys.argv[1] if len(sys.argv) > 1 else "Half_Marathon_Plan_2027.xlsx"
F = "Arial"
HDR = PatternFill("solid", fgColor="0F4C5C")
SEC = PatternFill("solid", fgColor="D9EEF2")
INP = PatternFill("solid", fgColor="FFF2B3")
GREEN = PatternFill("solid", fgColor="C9EBD2")
AMBER = PatternFill("solid", fgColor="FFE3A8")
RED = PatternFill("solid", fgColor="F6C2BE")
PH = {"Foundation": "E8F4EA", "Aerobic development": "E6EEF9", "Race-specific endurance": "FCEBDD",
      "Peak": "F8DCD2", "Taper": "F1E8F7", "Race week": "EDE3F7"}
TYPE_FONT = {"Run": "1B6E3A", "Swim": "1F5FA8", "Strength": "9A4A00", "Mobility": "6B3FA0", "Rest": "666666"}
thin = Side(style="thin", color="C8D3D6")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="top", wrap_text=True)


def hdr(ws, row, vals, widths=None, col=1):
    for i, v in enumerate(vals):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = B
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(col + i)].width = w


def title(ws, t, sub=None):
    ws["A1"] = t
    ws["A1"].font = Font(name=F, bold=True, size=15, color="0F4C5C")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=F, italic=True, size=10, color="555555")


def section(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = Font(name=F, bold=True, size=12, color="0F4C5C")
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = SEC


def cell(ws, r, c, v, inp=False, bold=False, fmt=None, color=None, align=WRAP):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name=F, size=10, bold=bold, color=color or ("0000FF" if inp and not str(v).startswith("=") and v not in (None, "") else "000000"))
    x.alignment = align
    x.border = B
    if inp:
        x.fill = INP
    if fmt:
        x.number_format = fmt
    return x


def rows(ws, start, data, inputs=()):
    for r, vals in enumerate(data, start):
        for c, v in enumerate(vals, 1):
            cell(ws, r, c, v, inp=c in inputs)
    return start + len(data)


sessions = P.build_sessions()
days = []
d = P.START
while d <= P.RACE:
    days.append(d)
    d += dt.timedelta(days=1)

wb = Workbook()

# ======================= Overview =======================
ov = wb.active
ov.title = "Overview"
title(ov, "Kilimanjaro Half Marathon Plan — 29 Oct 2026 to 7 Feb 2027",
      "Yellow cells are yours to fill in. Everything else updates itself. Full explanations are in the plan document.")
for col, w in zip("ABCDEF", [34, 22, 18, 18, 18, 30]):
    ov.column_dimensions[col].width = w

section(ov, 4, "Your numbers", 6)
inputs = [
    ("Start weight (kg)", 87, "You, 29 Oct 2026"),
    ("Height (m)", 1.62, "You"),
    ("Age", 26, "You, earlier"),
    ("Activity factor before training", 1.2, "Desk job (assumption)"),
    ("Race date", P.RACE, "Kilimanjaro Half Marathon, Moshi, 07:00 start"),
    ("Race cut-off (clock time)", "12:30", "Organiser's race info page"),
]
for i, (lab, val, note) in enumerate(inputs):
    r = 5 + i
    cell(ov, r, 1, lab, bold=True)
    cell(ov, r, 2, val, inp=True, fmt="dddd d mmm yyyy" if isinstance(val, dt.date) else None)
    cell(ov, r, 3, note, color="555555")
    ov.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)

section(ov, 12, "Calculated", 6)
calc = [
    ("BMI", "=B5/B6^2", "0.0", "Screening number only; waist tells you more"),
    ("Resting energy, Mifflin–St Jeor (kcal)", "=10*B5+6.25*B6*100-5*B7+5", "#,##0", "Starting estimate, about ±10%"),
    ("Needs before training (kcal)", "=B14*B8", "#,##0", "Training adds about 250–450 kcal a day"),
    ("Pace needed for 3:00 (min per km)", "=180/21.0975/1440", "[m]:ss", "A maths reference, not your current pace"),
    ("Today's projection from a 60-min 5 km", "=60*(21.0975/5)^1.06*1.08/1440", "[h]:mm", "Riegel × 1.08 for hills, heat, beginner"),
]
for i, (lab, f, fmt, note) in enumerate(calc):
    r = 13 + i
    cell(ov, r, 1, lab, bold=True)
    cell(ov, r, 2, f, fmt=fmt)
    cell(ov, r, 3, note, color="555555")
    ov.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)

section(ov, 19, "Progress dashboard (updates from the Daily Plan, Daily Log and Weekly Review sheets)", 6)
dash = [
    ("Days to race", "=MAX(0,B9-TODAY())", "0"),
    ("Latest 7-day average weight (kg)", "=IFERROR(LOOKUP(2,1/ISNUMBER('Daily Log'!F5:F106),'Daily Log'!F5:F106),\"Log your weight\")", "0.0"),
    ("Change from start (kg)", "=IFERROR(B21-B5,\"\")", "+0.0;-0.0;0.0"),
    ("Sessions done", "=COUNTIF('Daily Plan'!K:K,\"Yes\")+COUNTIF('Daily Plan'!K:K,\"Partial\")", "0"),
    ("Sessions planned (whole plan)", "=COUNTIFS('Daily Plan'!F:F,\"<>Rest\",'Daily Plan'!F:F,\"<>Type\",'Daily Plan'!F:F,\"<>\")", "0"),
    ("Completion", "=IFERROR(B23/B24,0)", "0%"),
    ("Longest run-walk logged (min)", "=_xlfn.MAXIFS('Daily Plan'!L:L,'Daily Plan'!F:F,\"Run\")", "0"),
    ("Longest continuous swim logged (m)", "=MAX('Weekly Review'!T:T)", "0"),
    ("Latest 5 km test (min)", "=IFERROR(LOOKUP(2,1/ISNUMBER('Weekly Review'!U5:U19),'Weekly Review'!U5:U19),\"Not tested yet\")", "0.0"),
    ("Race projection from latest test", "=IFERROR(B28*(21.0975/5)^1.06*1.08/1440,\"\")", "[h]:mm"),
    ("Suggested race plan", "=IFERROR(LOOKUP(B28,{0,36.01,40.01,43.01,47.01},{\"3:00\",\"3:15\",\"3:30\",\"3:45\",\"Finish strong (4:00)\"}),\"After your first test\")", None),
    ("This week's decision", "=IFERROR(INDEX('Weekly Review'!V5:V19,MATCH(MAX(0,MIN(14,INT((TODAY()-DATE(2026,11,2))/7)+1)),'Weekly Review'!A5:A19,0)),\"\")", None),
]
for i, (lab, f, fmt) in enumerate(dash):
    r = 20 + i
    cell(ov, r, 1, lab, bold=True)
    c = cell(ov, r, 2, f, bold=True, fmt=fmt)
ov["B28"].number_format = "0.0"
for code, fill in (("Green", GREEN), ("Amber", AMBER), ("Red", RED)):
    ov.conditional_formatting.add("B31", FormulaRule(formula=[f'LEFT($B$31,{len(code)})="{code}"'], fill=fill))

section(ov, 33, "Phases", 6)
hdr(ov, 34, ["Phase", "Weeks", "Dates", "Focus"])
ov.merge_cells("D34:F34")
phases = [
    ("Foundation and movement adaptation", "0–4", "29 Oct–29 Nov", "Short run intervals, strength habits, first swims, first 5 km test"),
    ("Aerobic development", "5–8", "30 Nov–27 Dec", "Longer run intervals, strides, first hill walks, second test"),
    ("Race-specific endurance", "9–11", "28 Dec–17 Jan", "Steady blocks, hill repeats, long sessions to 2 hours, deciding test"),
    ("Peak", "12", "18–24 Jan", "2 h 30 min race rehearsal with race fuelling"),
    ("Taper and race", "13–14", "25 Jan–7 Feb", "Less volume, same rhythm, travel, race"),
]
for i, p in enumerate(phases):
    r = 35 + i
    for c, v in enumerate(p, 1):
        cell(ov, r, c, v)
    ov.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)

section(ov, 41, "How to use this workbook", 6)
howto = [
    ("Daily Plan", "Every session, 29 Oct–7 Feb, written out in full. After each one fill the yellow cells: Done (Yes, Partial, No), minutes, distance, RPE 1–10, notes."),
    ("Example Daily Plan entry", "Done: Yes · Minutes: 32 · Distance: 3.1 · RPE: 4 · Notes: humid, extra walk break at 20 min"),
    ("Daily Log", "Every morning: weight, resting pulse, sleep, soreness, pain, energy, mood, motivation, nutrition and fasting scores. The 7-day average weight calculates itself."),
    ("Example Daily Log entry", "Weight 86.6 · Pulse 64 · Sleep 7.0 h, quality 4 · Soreness 2 · Pain 0 · Energy 4 · Mood 4 · Motivation 4 · Nutrition 4"),
    ("Weekly Review", "Sunday evening: add waist, swim distance and any 5 km test. The Decision column turns green (progress), amber (repeat the week) or red (step back)."),
    ("Nutrition, Meal Plan, Foods", "Targets by day type, a 7-day meal plan whose totals are formulas from the Foods table. Change a quantity or a food and the totals update."),
    ("Race Plan", "Pacing plans and checkpoints for each finish time, calculated from the climb and descent factors."),
]
rows(ov, 42, howto)
for r in range(42, 49):
    ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ov.row_dimensions[r].height = 30

# ======================= Daily Plan =======================
dp = wb.create_sheet("Daily Plan")
title(dp, "Daily Plan — every session, 29 October 2026 to 7 February 2027",
      "Fill the yellow columns after each session. Done: Yes / Partial / No. Minutes = total time. Distance in km for runs, m for swims.")
hdr(dp, 4, ["Date", "Week", "Day", "Phase", "Time", "Type", "Session", "Full details", "Planned", "Unit",
            "Done", "Minutes", "Distance", "RPE (1–10)", "Notes"],
    [12, 6, 10, 12, 13, 9, 30, 95, 8, 7, 8, 8, 9, 8, 26])
dv_done = DataValidation(type="list", formula1='"Yes,Partial,No"', allow_blank=True)
dv_rpe = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True)
dp.add_data_validation(dv_done)
dp.add_data_validation(dv_rpe)
for i, s in enumerate(sessions):
    r = 5 + i
    vals = [dt.date.fromisoformat(s["date"]), s["week"], s["day"], s["phase"], s["time_slot"], s["type"], s["title"],
            s["details"], s["planned"], s["unit"]]
    for c, v in enumerate(vals, 1):
        x = cell(dp, r, c, v, bold=(c == 7), color=TYPE_FONT.get(s["type"]) if c in (6, 7) else None)
        x.fill = PatternFill("solid", fgColor=PH[s["phase"]])
    dp.cell(row=r, column=1).number_format = "ddd d mmm"
    for c in range(11, 16):
        cell(dp, r, c, None, inp=True)
    dv_done.add(f"K{r}")
    dv_rpe.add(f"N{r}")
LAST_DP = 4 + len(sessions)
dp.freeze_panes = "C5"
dp.auto_filter.ref = f"A4:O{LAST_DP}"

# ======================= Daily Log =======================
dl = wb.create_sheet("Daily Log")
title(dl, "Daily Log — two minutes every morning",
      "Weight after the toilet, before food or drink. Scores: soreness and pain 0–10; others 1–5 (5 = best). Fasting tolerance only on Mon and Thu.")
hdr(dl, 4, ["Date", "Week", "Day", "Fast day", "Weight (kg)", "7-day avg (kg)", "Resting pulse", "Sleep (h)",
            "Sleep quality 1–5", "Soreness 0–10", "Pain 0–10", "Pain where", "Energy 1–5", "Mood 1–5", "Motivation 1–5",
            "Nutrition 1–5", "Fasting tolerance 1–5", "Warning signs (Yes/No)", "Notes"],
    [12, 6, 10, 8, 9, 9, 9, 8, 9, 9, 8, 14, 8, 8, 9, 9, 10, 10, 26])
dv15 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
dv010 = DataValidation(type="whole", operator="between", formula1="0", formula2="10", allow_blank=True)
dvyn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
for v in (dv15, dv010, dvyn):
    dl.add_data_validation(v)
for i, d in enumerate(days):
    r = 5 + i
    fast = d.weekday() in (0, 3)
    w = P.week_of(d)
    for c, v in enumerate([d, w, d.strftime("%A"), "Yes" if fast else ""], 1):
        x = cell(dl, r, c, v)
        x.fill = PatternFill("solid", fgColor=PH[P.PHASES[w]])
    dl.cell(row=r, column=1).number_format = "ddd d mmm"
    top = max(5, r - 6)
    cell(dl, r, 6, f'=IF(COUNT(E{top}:E{r})=0,"",AVERAGE(E{top}:E{r}))', fmt="0.0")
    for c in [5] + list(range(7, 20)):
        if c == 17 and not fast:
            x = cell(dl, r, c, "—")
            x.fill = PatternFill("solid", fgColor="EEEEEE")
            x.alignment = CEN
            continue
        cell(dl, r, c, None, inp=True)
    dl.cell(row=r, column=5).number_format = "0.0"
    for col in "IMNOPQ":
        dv15.add(f"{col}{r}")
    for col in "JK":
        dv010.add(f"{col}{r}")
    dvyn.add(f"R{r}")
LAST_DL = 4 + len(days)
dl.freeze_panes = "E5"
dl.conditional_formatting.add(f"K5:K{LAST_DL}", FormulaRule(formula=["AND(ISNUMBER(K5),K5>=4)"], fill=RED))
dl.conditional_formatting.add(f"R5:R{LAST_DL}", FormulaRule(formula=['R5="Yes"'], fill=RED))

# ======================= Weekly Review =======================
wr = wb.create_sheet("Weekly Review")
title(wr, "Weekly Review — the Sunday decision",
      "Fill waist, swim and any 5 km test (yellow). Decision: Green = progress, Amber = repeat the week, Red = step back (plan section 14).")
cols = ["Week", "Starts", "Phase", "Sessions planned", "Sessions done", "Completion", "Run min planned", "Run min done",
        "Long session planned (min)", "Long session done (min)", "Avg run RPE", "Avg sleep (h)", "Avg sleep quality",
        "Max pain", "Avg energy", "Avg mood", "Avg weight (kg)", "Change (kg)", "Waist (cm)", "Swim continuous (m)",
        "5 km test (min)", "Decision", "Notes"]
hdr(wr, 4, cols, [6, 11, 14, 9, 9, 10, 9, 9, 10, 10, 8, 8, 9, 7, 8, 8, 9, 8, 8, 10, 9, 30, 26])
DPr = f"'Daily Plan'!$B$5:$B${LAST_DP}"
DLr = f"'Daily Log'!$B$5:$B${LAST_DL}"
for w in range(15):
    r = 5 + w
    start = P.START if w == 0 else P.WEEK1 + dt.timedelta(weeks=w - 1)
    t, f, s = P.RUN[w]
    vals = {
        1: w, 2: start, 3: P.PHASES[w],
        4: f"=COUNTIFS({DPr},A{r},'Daily Plan'!$F$5:$F${LAST_DP},\"<>Rest\")",
        5: f"=COUNTIFS({DPr},A{r},'Daily Plan'!$K$5:$K${LAST_DP},\"Yes\")+COUNTIFS({DPr},A{r},'Daily Plan'!$K$5:$K${LAST_DP},\"Partial\")",
        6: f"=IFERROR(E{r}/D{r},\"\")",
        7: f"=SUMIFS('Daily Plan'!$I$5:$I${LAST_DP},{DPr},A{r},'Daily Plan'!$F$5:$F${LAST_DP},\"Run\",'Daily Plan'!$J$5:$J${LAST_DP},\"min\")",
        8: f"=SUMIFS('Daily Plan'!$L$5:$L${LAST_DP},{DPr},A{r},'Daily Plan'!$F$5:$F${LAST_DP},\"Run\")",
        9: s or "",
        10: f"=IFERROR(INDEX('Daily Plan'!$L$5:$L${LAST_DP},MATCH(1,INDEX(({DPr}=A{r})*('Daily Plan'!$C$5:$C${LAST_DP}=\"Saturday\")*('Daily Plan'!$F$5:$F${LAST_DP}=\"Run\"),0),0)),\"\")",
        11: f"=IFERROR(AVERAGEIFS('Daily Plan'!$N$5:$N${LAST_DP},{DPr},A{r},'Daily Plan'!$F$5:$F${LAST_DP},\"Run\"),\"\")",
        12: f"=IFERROR(AVERAGEIFS('Daily Log'!$H$5:$H${LAST_DL},{DLr},A{r}),\"\")",
        13: f"=IFERROR(AVERAGEIFS('Daily Log'!$I$5:$I${LAST_DL},{DLr},A{r}),\"\")",
        14: f"=IF(COUNTIFS({DLr},A{r},'Daily Log'!$K$5:$K${LAST_DL},\">=0\")=0,\"\",_xlfn.MAXIFS('Daily Log'!$K$5:$K${LAST_DL},{DLr},A{r}))",
        15: f"=IFERROR(AVERAGEIFS('Daily Log'!$M$5:$M${LAST_DL},{DLr},A{r}),\"\")",
        16: f"=IFERROR(AVERAGEIFS('Daily Log'!$N$5:$N${LAST_DL},{DLr},A{r}),\"\")",
        17: f"=IFERROR(AVERAGEIFS('Daily Log'!$E$5:$E${LAST_DL},{DLr},A{r}),\"\")",
        18: "" if w == 0 else f"=IF(OR(Q{r}=\"\",Q{r-1}=\"\"),\"\",Q{r}-Q{r-1})",
        22: (f"=IF(E{r}=0,\"\",IF(OR(AND(N{r}<>\"\",N{r}>=4),AND(F{r}<>\"\",F{r}<0.6),AND(L{r}<>\"\",L{r}<6)),"
             f"\"Red: step back\",IF(OR(AND(F{r}<>\"\",F{r}<0.85),AND(K{r}<>\"\",K{r}>=5),AND(N{r}<>\"\",N{r}>=3),"
             f"AND(L{r}<>\"\",L{r}<6.5),AND(O{r}<>\"\",O{r}<3),AND(P{r}<>\"\",P{r}<3)),\"Amber: repeat the week\",\"Green: progress\")))"),
    }
    for c in range(1, 24):
        v = vals.get(c)
        is_inp = c in (19, 20, 21, 23)
        x = cell(wr, r, c, v if v is not None else None, inp=is_inp, align=CEN if c != 23 else WRAP)
        if not is_inp:
            x.fill = PatternFill("solid", fgColor=PH[P.PHASES[w]])
    wr.cell(row=r, column=2).number_format = "d mmm"
    wr.cell(row=r, column=6).number_format = "0%"
    for c in (11, 12, 13, 15, 16):
        wr.cell(row=r, column=c).number_format = "0.0"
    wr.cell(row=r, column=17).number_format = "0.0"
    wr.cell(row=r, column=18).number_format = "+0.0;-0.0;0.0"
    wr.cell(row=r, column=21).number_format = "0.0"
    if w in (4, 8, 11):
        wr.cell(row=r, column=21).comment = Comment("5 km test week: enter your time in minutes, e.g. 47.5", "Plan")
for code, fill in (("Green", GREEN), ("Amber", AMBER), ("Red", RED)):
    wr.conditional_formatting.add("V5:V19", FormulaRule(formula=[f'LEFT($V5,{len(code)})="{code}"'], fill=fill))
wr.freeze_panes = "C5"
wr["A21"] = "Decision rules: Red if max pain ≥ 4, completion < 60% or sleep < 6 h. Amber if completion < 85%, run RPE ≥ 5, pain 3, sleep < 6.5 h, or energy or mood < 3. Otherwise Green. Read the pattern, not one number."
wr["A21"].font = Font(name=F, italic=True, size=9)
ch = LineChart()
ch.title = "Average weight by week (kg)"
ch.height, ch.width = 7.5, 16
ch.add_data(Reference(wr, min_col=17, min_row=4, max_row=19), titles_from_data=True)
ch.set_categories(Reference(wr, min_col=1, min_row=5, max_row=19))
ch.y_axis.title = "kg"
ch.x_axis.title = "Week"
wr.add_chart(ch, "C23")

# ======================= Week Plan =======================
wk = wb.create_sheet("Week Plan")
title(wk, "Fourteen-week overview", "Run time includes walking intervals. 1:2 = run 1 min, walk 2 min.")
hdr(wk, 4, ["Week", "Starts", "Phase", "Run:walk", "Tue (min)", "Fri (min)", "Sat long (min)", "Run time (min)",
            "Longest (≈ km)", "Strength", "Swims (optional)", "Objective"], [6, 11, 18, 8, 8, 8, 9, 9, 9, 10, 30, 60])
for w in range(15):
    r = 5 + w
    t, f, s = P.RUN[w]
    start = P.START if w == 0 else P.WEEK1 + dt.timedelta(weeks=w - 1)
    sw = []
    if P.SWIM_WED.get(w):
        sw.append(P.SWIM_WED[w][0])
    if P.SWIM_SUN_CONT.get(w) is not None:
        sw.append(f"{P.SWIM_SUN_CONT[w]} m continuous")
    strength = {0: 0, 13: 1, 14: 1}.get(w, 2)
    obj = P.WEEK_GOAL[w] + (" " + P.RECOVERY[w] if w in P.RECOVERY else "")
    tv = P.TT_MIN if t == "TT" else (t or 0)
    fv = P.TT_MIN if f == "TT" else (f or 0)
    vals = [w, start, P.PHASES[w], f"{P.RATIO[w][0]}:{P.RATIO[w][1]}", "5 km test" if t == "TT" else tv,
            "5 km test" if f == "TT" else fv, s or "Race", f"={tv}+{fv}+{s or 0}",
            round((s or 0) / P.EST_PACE[w]) if s else 21.1, strength, "; ".join(sw) or "—", obj]
    for c, v in enumerate(vals, 1):
        x = cell(wk, r, c, v)
        x.fill = PatternFill("solid", fgColor=PH[P.PHASES[w]])
    wk.cell(row=r, column=2).number_format = "d mmm"

# ======================= Strength =======================
st = wb.create_sheet("Strength")
title(st, "Strength A (Wednesday) and B (Tuesday evening)",
      "Stop each set with 2–3 reps left (RPE 7). Cutback weeks 4, 8, 11: 2 sets. Week 13: A only, 2 sets. Week 14: upper body only.")
hdr(st, 4, ["Session", "Exercise", "Sets × reps", "Rest", "Technique cues", "Progression"], [12, 26, 16, 9, 55, 45])
strength = [
    ("A", "Sit-to-stand → squat", "3 × 8–12", "60–90 s", "Feet hip-width, knees over toes, chest up, sit back", "Chair → no chair → backpack → 3 s lowering"),
    ("A", "Glute bridge", "3 × 10–12", "60 s", "Ribs down, push through heels, 2 s squeeze", "Two legs → single leg"),
    ("A", "Step-up (or reverse lunge)", "3 × 8 per leg", "60–90 s", "Safe 20–30 cm step; no step: reverse lunge holding a wall", "Add the backpack"),
    ("A", "Hip hinge", "3 × 10", "60 s", "Soft knees, flat back, hips back", "Good morning → single-leg RDL (hand on wall) → with backpack"),
    ("A", "Calf raise", "3 × 12–15", "45–60 s", "Step edge, full range, 3 s down", "Two legs → single leg → single leg + backpack"),
    ("A", "Bent-knee calf raise", "2 × 15", "45 s", "Knees bent 30°", "Single leg"),
    ("A", "Dead bug", "3 × 8 per side", "45 s", "Lower back pressed down", "Slower, straighter legs"),
    ("A", "Side plank", "3 × 20–40 s per side", "45 s", "Straight line ear-hip-ankle", "Knees → feet → top-leg lift"),
    ("A", "Pelvic-floor finisher", "1 × 10", "—", "Glute bridge with 5 s pelvic-floor squeeze", "Hold 8 s"),
    ("B", "Push-up", "3 × 8–12", "60 s", "Body straight, elbows about 45°", "Wall → table → knees → full"),
    ("B", "Backpack bent-over row", "3 × 10–12", "60 s", "Hinge, flat back, bag to lower ribs", "Heavier bag, 2 s squeeze"),
    ("B", "Overhead press / pike push-up", "3 × 8–10", "60 s", "Ribs down, press straight up", "Heavier bag; feet raised"),
    ("B", "Single-leg balance", "3 × 30 s per leg", "30 s", "Soft knee, hips level", "Eyes closed → folded towel"),
    ("B", "Single-leg calf lowering", "2 × 10 per leg", "45 s", "Up on two, down on one over 3 s", "Add backpack"),
    ("B", "Towel scrunch + tibialis raise", "2 × 20 + 2 × 15", "30 s", "Toes scrunch a towel; back to wall, lift toes", "More reps, slower"),
    ("B", "Bird dog", "3 × 8 per side", "45 s", "Hips level, reach long", "Hold 3 s"),
    ("B", "Front plank", "3 × 20–45 s", "45 s", "Squeeze glutes, no sag", "Longer → shoulder taps"),
    ("B", "Pelvic-floor finisher", "10 squeezes", "—", "5 s squeeze in a deep squat hold", "Hold 8 s"),
]
rows(st, 5, strength)
section(st, 25, "Soreness or injury?", 6)
rows(st, 26, [
    ("Normal", "Dull, spread across a muscle, peaks 24–48 h after, eases as you warm up — train as planned"),
    ("Warning", "Sharp or pinpoint pain, pain on a bone, swelling, limping, pain at night, pain lasting 3+ days"),
    ("Rule", "Pain up to 3/10 that stays stable or eases: OK. Above 3 or rising: stop, 3 days walking/swimming only, physio if it lasts a week"),
])
for r in range(26, 29):
    st.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

# ======================= Swim =======================
sw = wb.create_sheet("Swim Plan")
title(sw, "Swim plan (optional) — from 100 m to 600–700 m continuous",
      "Wednesday repeats at RPE 4–5; Sunday continuous very easy (RPE 3–4). Never on a run day or a fast day.")
hdr(sw, 4, ["Week", "Wednesday main set", "Rest", "Wednesday total (m)", "Sunday continuous (m)", "Sunday total (m)", "Your longest non-stop (m)"],
    [6, 22, 8, 12, 12, 12, 14])
for w in range(15):
    r = 5 + w
    wed = P.SWIM_WED.get(w)
    sun = P.SWIM_SUN_CONT.get(w)
    vals = [w, wed[0] if wed else "—", f"{wed[3]} s" if wed else "", f"=100+100+{wed[1]}*{wed[2]}+100" if wed else "",
            sun if sun is not None else "Race", f"=100+100+{sun}+200+100" if sun is not None else "",
            f"=IFERROR(INDEX('Weekly Review'!T5:T19,MATCH(A{r},'Weekly Review'!A5:A19,0)),\"\")"]
    for c, v in enumerate(vals, 1):
        cell(sw, r, c, v)
sw["A21"] = ("Every swim: warm-up 4 × 25 m easy; drills 4 × 25 m (side kick, catch-up, fingertip drag, six-kick switch); "
             "main set; cool-down 50–100 m easy. Sunday adds 4 × 50 m relaxed after the continuous swim.")
sw["A21"].font = Font(name=F, italic=True, size=9)

# ======================= Foods =======================
fo = wb.create_sheet("Foods")
title(fo, "Food values", "Typical values for the closest standard food; local recipes vary about ±15%. Edit to match your foods.")
hdr(fo, 4, ["Key", "Food", "Basis", "Divide by", "kcal", "Protein (g)", "Carbs (g)", "Fat (g)"], [14, 30, 12, 9, 8, 9, 9, 8])
food_rows = []
for k, v in M.F.items():
    food_rows.append((k, M.LABEL[k], "per 100 g", 100) + v)
for k, v in M.U.items():
    food_rows.append((k, M.LABEL[k], "per unit", 1) + v)
for i, fr in enumerate(food_rows):
    for c, v in enumerate(fr, 1):
        cell(fo, 5 + i, c, v, inp=c >= 5)
LAST_FO = 4 + len(food_rows)
FK = f"Foods!$A$5:$A${LAST_FO}"

# ======================= Meal Plan =======================
mp = wb.create_sheet("Meal Plan")
title(mp, "Seven-day Tanzanian meal plan", "Change a quantity (yellow) or a food key and the totals update. Keys are listed on the Foods sheet.")
hdr(mp, 4, ["Day", "Time", "Meal", "Food key", "Quantity", "Unit", "kcal", "Protein (g)", "Carbs (g)", "Fat (g)", "Note"],
    [26, 7, 22, 14, 9, 9, 8, 9, 9, 8, 50])
r = 5
for dname, day in M.DAYS.items():
    for time, meal, foods, note in day["meals"]:
        first = True
        for fkey, q in foods:
            unit = "g" if fkey in M.F else "unit"
            vals = [dname, time, meal, fkey, q, unit,
                    f"=E{r}/INDEX(Foods!$D$5:$D${LAST_FO},MATCH(D{r},{FK},0))*INDEX(Foods!$E$5:$E${LAST_FO},MATCH(D{r},{FK},0))",
                    f"=E{r}/INDEX(Foods!$D$5:$D${LAST_FO},MATCH(D{r},{FK},0))*INDEX(Foods!$F$5:$F${LAST_FO},MATCH(D{r},{FK},0))",
                    f"=E{r}/INDEX(Foods!$D$5:$D${LAST_FO},MATCH(D{r},{FK},0))*INDEX(Foods!$G$5:$G${LAST_FO},MATCH(D{r},{FK},0))",
                    f"=E{r}/INDEX(Foods!$D$5:$D${LAST_FO},MATCH(D{r},{FK},0))*INDEX(Foods!$H$5:$H${LAST_FO},MATCH(D{r},{FK},0))",
                    note if first else ""]
            for c, v in enumerate(vals, 1):
                cell(mp, r, c, v, inp=c in (4, 5))
            for c in range(7, 11):
                mp.cell(row=r, column=c).number_format = "0"
            first = False
            r += 1
LAST_MP = r - 1
dv_food = DataValidation(type="list", formula1=f"=Foods!$A$5:$A${LAST_FO}", allow_blank=False)
mp.add_data_validation(dv_food)
dv_food.add(f"D5:D{LAST_MP}")
mp.freeze_panes = "A5"

# ======================= Nutrition =======================
nu = wb.create_sheet("Nutrition")
title(nu, "Daily targets and meal-plan totals", "Totals are calculated from the Meal Plan sheet. Recalculate every 2 weeks from your weight trend.")
hdr(nu, 4, ["Day type", "Target kcal", "Target protein (g)", "Target carbs (g)", "Target fat (g)", "Meal plan kcal",
            "Meal plan protein", "Meal plan carbs", "Meal plan fat", "kcal difference"], [30, 10, 10, 10, 10, 10, 10, 10, 10, 10])
for i, (dname, day) in enumerate(M.DAYS.items()):
    r = 5 + i
    t = day["target"]
    vals = [dname, t[0], t[1], t[2], t[3]]
    for c, v in enumerate(vals, 1):
        cell(nu, r, c, v, inp=c >= 2)
    for j, col in enumerate("GHIJ"):
        cell(nu, r, 6 + j, f"=SUMIFS('Meal Plan'!${col}$5:${col}${LAST_MP},'Meal Plan'!$A$5:$A${LAST_MP},A{r})", fmt="#,##0")
    cell(nu, r, 10, f"=F{r}-B{r}", fmt="+#,##0;-#,##0;0")
cell(nu, 12, 1, "Weekly average", bold=True)
for c in range(2, 11):
    col = get_column_letter(c)
    cell(nu, 12, c, f"=AVERAGE({col}5:{col}11)", bold=True, fmt="#,##0")
section(nu, 14, "Every two weeks: compare your 7-day average weight with two weeks earlier", 10)
rows(nu, 15, [
    ("Lost more than 1.6 kg in 2 weeks", "Add 200 kcal a day, mostly carbohydrate around runs"),
    ("Lost 0.6–1.6 kg", "On track: change nothing"),
    ("Lost under 0.6 kg, training and sleep good", "Remove 150 kcal a day from oil, sugar or snacks — never protein or pre-run food"),
    ("Lost under 0.6 kg, but tired or sleeping badly", "Keep calories; fix sleep; reassess in 2 weeks"),
    ("From week 9", "Add about 150 kcal of carbohydrate on Tue and Fri, 250 kcal on Sat"),
    ("Race week (Fri, Sat)", "Carbohydrate about 4–5 g/kg (350–430 g), low fibre and fat"),
])
for rr in range(15, 21):
    nu.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=10)

# ======================= Race Plan =======================
rp = wb.create_sheet("Race Plan")
title(rp, "Race pacing — Kilimanjaro Half Marathon, 07:00 start, 12:30 cut-off",
      "The first 8 km climb about 290 m. Factors below are planning assumptions; run by effort. Change the yellow cells to test other plans.")
cell(rp, 4, 1, "Climb factor (0–8 km)", bold=True); cell(rp, 4, 2, 1.15, inp=True)
cell(rp, 5, 1, "Dirt section factor (8–10 km)", bold=True); cell(rp, 5, 2, 1.0, inp=True)
cell(rp, 6, 1, "Return factor (10–21.1 km, calculated)", bold=True); cell(rp, 6, 2, "=(21.0975-8*B4-2*B5)/11.0975", fmt="0.000")
cell(rp, 7, 1, "Hill, heat and beginner allowance", bold=True); cell(rp, 7, 2, 1.08, inp=True)
hdr(rp, 9, ["Plan (finish, h:mm)", "Use it if your 5 km test is", "Average pace", "Climb pace", "Dirt pace", "Return pace",
            "5 km", "10 km", "15 km", "20 km", "Finish clock time"], [26, 20, 11, 11, 11, 11, 9, 9, 9, 9, 12])
plans = [("3:00", 180, "36 min or faster"), ("3:15", 195, "37–40 min"), ("3:30", 210, "41–43 min"),
         ("3:45", 225, "44–47 min"), ("Finish strong (4:00)", 240, "48 min or slower")]
for i, (name, mins_, cond) in enumerate(plans):
    r = 10 + i
    cell(rp, r, 1, name, bold=True)
    cell(rp, r, 2, cond)
    cell(rp, r, 12, mins_, inp=True)
    p = f"($L{r}/21.0975)"

    def at(km):
        parts = [f"MIN({km},8)*$B$4*{p}"]
        if km > 8:
            parts.append(f"MIN({km}-8,2)*$B$5*{p}")
        if km > 10:
            parts.append(f"({km}-10)*$B$6*{p}")
        return "+".join(parts)
    cell(rp, r, 3, f"={p}/1440", fmt="[m]:ss")
    cell(rp, r, 4, f"={p}*$B$4/1440", fmt="[m]:ss")
    cell(rp, r, 5, f"={p}*$B$5/1440", fmt="[m]:ss")
    cell(rp, r, 6, f"={p}*$B$6/1440", fmt="[m]:ss")
    for j, km in enumerate((5, 10, 15, 20)):
        cell(rp, r, 7 + j, f"=({at(km)})/1440", fmt="[h]:mm")
    cell(rp, r, 11, f"=TIME(7,0,0)+L{r}/1440", fmt="hh:mm")
rp.cell(row=9, column=12, value="Finish (min)").font = Font(name=F, bold=True, color="FFFFFF", size=10)
rp.cell(row=9, column=12).fill = HDR
rp.column_dimensions["L"].width = 10
section(rp, 16, "Your 5 km test → projection", 11)
cell(rp, 17, 1, "5 km test time (min)", bold=True); cell(rp, 17, 2, 45, inp=True)
cell(rp, 18, 1, "Flat-course projection", bold=True); cell(rp, 18, 2, "=B17*(21.0975/5)^1.06/1440", fmt="[h]:mm")
cell(rp, 19, 1, "This course (× allowance)", bold=True); cell(rp, 19, 2, "=B17*(21.0975/5)^1.06*B7/1440", fmt="[h]:mm")
cell(rp, 20, 1, "Plan to use", bold=True)
cell(rp, 20, 2, '=LOOKUP(B17,{0,36.01,40.01,43.01,47.01},{"3:00","3:15","3:30","3:45","Finish strong (4:00)"})')
section(rp, 22, "Race morning", 11)
rows(rp, 23, [
    ("03:45", "Wake, toilet, water"),
    ("04:00", "Breakfast: 2–3 slices white bread + honey + 1 tbsp peanut butter, 1 banana, sweet tea or uji (about 100–120 g carbohydrate)"),
    ("05:30", "Leave for the start; sip water to thirst"),
    ("06:30", "Warm-up: 10 min easy walk, leg swings, ankle rocks"),
    ("06:50", "Line up in the back half of the field"),
    ("07:00 →", "Run-walk from km 1; walk the steep climb; carbohydrate every 20–25 min; drink to thirst at every water point"),
])
for rr in range(23, 29):
    rp.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=11)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT, "sessions", len(sessions), "days", len(days), "meal rows", LAST_MP - 4)
